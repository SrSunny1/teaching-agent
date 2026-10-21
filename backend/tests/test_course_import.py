"""Offline checks for course imports, provenance and labor supplements."""
import asyncio
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from knowledge.ingest import RAW, chunk, read_docx
from app.tools.search_knowledge import _run


class CourseImportTests(unittest.TestCase):
    def test_docx_tables_preserve_order(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'table.docx'
            doc = Document()
            doc.add_paragraph('before')
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = 'material'
            table.cell(0, 1).text = 'hardness'
            table.cell(1, 0).text = 'steel'
            table.cell(1, 1).text = 'value'
            doc.add_paragraph('after')
            doc.save(path)
            self.assertEqual(read_docx(path), ['before', 'material | hardness', 'steel | value', 'after'])

    def test_task_table_is_not_a_module_heading(self):
        rows = chunk(['模块一：安全巡查', '教学内容说明需要足够长度以形成一个独立可检索的段落。',
                      '【原文段落10】四、项目预期成效、标志性成果和绩效评价指标',
                      '【原文段落11】模块一：安全巡查\t2学时/次\t4轮\t约40人'], '任务书')
        self.assertEqual(rows[-1]['metadata']['experiment'], '劳育课程总览')
        self.assertIn('2学时/次', rows[-1]['text'])

    def test_manifest_and_sources(self):
        manifest = json.loads((RAW.parent / 'import-manifest-2026-10-21.json').read_text(encoding='utf-8'))
        self.assertEqual(len(manifest['files']), 11)
        self.assertEqual(sum(f.get('pages', 0) for f in manifest['files']), 86)
        for entry in manifest['files']:
            self.assertTrue((RAW / entry['output']).is_file(), entry['name'])
        rows = [json.loads(line) for line in (ROOT.parent / 'data/chunks.jsonl').read_text(encoding='utf-8').splitlines()]
        sources = {r['metadata']['doc'] for r in rows}
        for entry in manifest['files'][1:]:
            self.assertTrue(any(Path(entry['output']).stem[3:] in name for name in sources), entry['name'])
        public = [r for r in rows if r['metadata']['source'] == '劳育公开补充']
        self.assertEqual(len(public), 5)
        for row in public:
            self.assertIn('https://', row['text'])
            self.assertIn('待教师确认', row['text'])
        task = '\n'.join(r['text'] for r in rows if r['metadata']['source'] == '任务书')
        self.assertNotRegex(task, r'(?<!\d)1[3-9]\d{9}(?!\d)|[\w.+-]+@[\w.-]+\.[a-zA-Z]+')
        self.assertIn('SiO₂', task)
        self.assertIn('800℃', task)

    def test_labor_retrieval_has_public_sources(self):
        for query in ['实验室安全巡查包括哪些内容', '材料制备为什么也属于劳动教育',
                      '仪器发生故障怎么办', '实验废弃物怎么分类收集', '劳动教育怎样记录与评价']:
            result = asyncio.run(_run({'query': query}))
            self.assertIsNotNone(result.artifact, query)
            self.assertTrue(any(c['source'] == '劳育公开补充' for c in result.artifact['chunks']), query)


if __name__ == '__main__':
    unittest.main()
